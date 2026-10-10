from __future__ import annotations

import base64
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.api_client import _user_content
from app.prompting import build_prompt, parse_answers, remap_to_original
from app.runner import BatchRunner, RunnerConfig, ScaleRunConfig
from app.scale_loader import (
    ScaleFile,
    ScaleImage,
    ScaleSheet,
    load_scale_file,
    paired_shuffle_compatible,
)
from app.webapp import app


# A 1 x 1 PNG used only to verify the Excel -> API image path without a network call.
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
    "VY5t6QAAAABJRU5ErkJggg=="
)


class ScalePipelineTests(unittest.TestCase):
    def test_single_standard_language_sheet_is_not_reused_for_other_language(self):
        with tempfile.TemporaryDirectory() as folder:
            for name, language in (("Sheet1", "en"), ("Sheet2", "ch")):
                with self.subTest(sheet=name):
                    path = Path(folder) / f"{name}.xlsx"
                    workbook = Workbook()
                    workbook.active.title = name
                    for value in ("Fixture", "Answer the question.", "1. Question."):
                        workbook.active.append([value])
                    workbook.save(path)
                    workbook.close()
                    scale = load_scale_file(path)
                    self.assertIsNotNone(getattr(scale, language))
                    self.assertIsNone(getattr(scale, "ch" if language == "en" else "en"))

    def test_nonstandard_sheet_fallback_does_not_duplicate_named_language(self):
        with tempfile.TemporaryDirectory() as folder:
            for names, expected in ((("Sheet2", "English"), ("English", "Sheet2")),
                                    (("First", "Second"), ("First", "Second"))):
                with self.subTest(sheets=names):
                    workbook = Workbook()
                    workbook.active.title = names[0]
                    workbook.create_sheet(names[1])
                    for sheet in workbook:
                        for value in (sheet.title, "Answer the question.", "1. Question."):
                            sheet.append([value])
                    path = Path(folder) / "fixture.xlsx"
                    workbook.save(path)
                    workbook.close()
                    scale = load_scale_file(path)
                    self.assertEqual((scale.en.title, scale.ch.title), expected)

    def test_loader_keeps_task_setting_as_context_and_extracts_image(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            png = tmp / "stimulus.png"
            png.write_bytes(PNG_1X1)
            xlsx = tmp / "dilemma.xlsx"

            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws["A1"] = "Dilemma"
            ws["A2"] = "Base scenario."
            ws["A3"] = "Task setting: You are Player 1."
            ws["A4"] = "1. Choose A or B."
            ws["A5"] = "2. Rate acceptability from 1 to 7."
            ws.add_image(ExcelImage(str(png)), "C3")
            ws_ch = wb.create_sheet("Sheet2")
            ws_ch["A1"] = "困境"
            ws_ch["A2"] = "基本情境。"
            ws_ch["A3"] = "任务设定：你是玩家1。"
            ws_ch["A4"] = "1. 请选择 A 或 B。"
            wb.save(xlsx)
            wb.close()

            scale = load_scale_file(xlsx)
            self.assertEqual(scale.en.n_items, 2)
            self.assertIn("Task setting", scale.en.instruction)
            self.assertNotIn("Task setting", scale.en.questions)
            self.assertEqual(len(scale.en.images), 1)
            self.assertTrue(scale.en.images[0].data_url.startswith("data:image/png;base64,"))
            self.assertEqual(scale.ch.n_items, 1)

    def test_strict_parser_never_grabs_numbers_from_prose(self):
        result = parse_answers("I choose B because 5 is greater than 1.", 2)
        self.assertEqual(result.status, "invalid_json")
        self.assertEqual([item.score for item in result.answers], [None, None])

    def test_strict_parser_keeps_uncertain_range_without_guessing(self):
        result = parse_answers(
            '{"answers":[{"display_index":1,"answer":"3-4"},'
            '{"display_index":2,"answer":"5"}]}',
            2,
        )
        self.assertEqual(result.status, "partial")
        self.assertEqual(result.answers[0].answer, "3-4")
        self.assertIsNone(result.answers[0].score)
        self.assertEqual(result.answers[0].parse_status, "range_unresolved")
        self.assertEqual(result.answers[1].score, 5.0)

    def test_null_answer_is_invalid_and_not_offered_as_a_template(self):
        sheet = ScaleSheet(
            language="ch",
            title="电车困境",
            instruction="中性情境。",
            questions=["选择 A 或 B", "给出 1 到 7 的评分"],
        )
        prompt = build_prompt(sheet, [0, 1], language="ch")
        self.assertNotIn('"answer":null', prompt)
        self.assertNotIn("严禁输出 null", prompt)

        raw = (
            '{"answers":[{"display_index":1,"answer":null},'
            '{"display_index":2,"answer":null}]}'
        )
        parsed = parse_answers(raw, 2)
        self.assertNotEqual(parsed.status, "ok")
        self.assertEqual([item.parse_status for item in parsed.answers], ["null_answer", "null_answer"])
        self.assertIn("Null answer", parsed.error)

    def test_structured_answers_and_order_remap(self):
        raw = (
            '{"answers":[{"display_index":1,"answer":"B"},'
            '{"display_index":2,"answer":7}]}'
        )
        parsed = parse_answers(raw, 2)
        self.assertEqual(parsed.status, "ok")
        original = remap_to_original(parsed.answers, [1, 0])
        self.assertEqual(original[0].score, 7.0)
        self.assertEqual(original[1].answer, "B")
        self.assertIsNone(original[1].score)

    def test_fixed_profile_prevents_obsolete_shuffle_override(self):
        runner = BatchRunner([], RunnerConfig())
        fixed = ScaleSheet(language="ch", title="Interview", instruction="", questions=["Question"], shuffle_supported=False)
        self.assertFalse(runner._should_shuffle("Interview", fixed))
        self.assertTrue(runner._should_shuffle("Independent items"))
        runner.config.scale_cfgs["Interview"] = ScaleRunConfig(scale_name="Interview", shuffle_items=True)
        self.assertFalse(runner._should_shuffle("Interview", fixed))

    def test_paired_shuffle_reuses_the_same_seed_for_both_languages(self):
        en = ScaleSheet(
            language="en",
            title="Fixture",
            instruction="Answer the items.",
            questions=["Item 1", "Item 2", "Item 3"],
            shuffle_supported=True,
        )
        ch = ScaleSheet(
            language="ch",
            title="测试量表",
            instruction="回答题目。",
            questions=["题目 1", "题目 2", "题目 3"],
            shuffle_supported=True,
        )
        scale = ScaleFile(path=Path("fixture.xlsx"), name="fixture", en=en, ch=ch)
        config = RunnerConfig(order_strategy="paired", random_seed=20260911)
        config.scale_cfgs[scale.name] = ScaleRunConfig(
            scale_name=scale.name, shuffle_items=True
        )
        runner = BatchRunner([scale], config)

        self.assertTrue(paired_shuffle_compatible(en, ch))
        self.assertTrue(runner._should_shuffle(scale.name, en))
        seed_en = runner._trial_seed(scale.name, "en", 1, True)
        seed_ch = runner._trial_seed(scale.name, "ch", 1, True)
        self.assertEqual(seed_en, seed_ch)
        self.assertNotEqual(seed_en, runner._trial_seed(scale.name, "en", 2, True))

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            scale.path = root / "fixture.xlsx"
            runner.config.results_root = root / "results"
            raw = json.dumps(
                {
                    "answers": [
                        {"display_index": index, "answer": f"choice-{index}"}
                        for index in range(1, 4)
                    ]
                }
            )
            with patch("app.runner.call_api", return_value=raw):
                en_path = runner._one_call(scale, en, "en", 1, "deepseek", "deepseek-v4-pro")
                ch_path = runner._one_call(scale, ch, "ch", 1, "deepseek", "deepseek-v4-pro")
            en_record = json.loads(en_path.read_text(encoding="utf-8"))
            ch_record = json.loads(ch_path.read_text(encoding="utf-8"))
            self.assertEqual(en_record["pair_id"], 1)
            self.assertEqual(ch_record["pair_id"], 1)
            self.assertEqual(en_record["order_strategy"], "paired")
            self.assertEqual(ch_record["order_strategy"], "paired")
            self.assertEqual(en_record["paired_trial_key"], ch_record["paired_trial_key"])
            self.assertEqual(en_record["shuffle_seed"], ch_record["shuffle_seed"])
            self.assertEqual(en_record["shuffle_order"], ch_record["shuffle_order"])

        runner.config.order_strategy = "fixed"
        self.assertFalse(runner._should_shuffle(scale.name, en))

    def test_paired_shuffle_requires_matching_bilingual_layouts(self):
        en = ScaleSheet(
            language="en", title="Fixture", instruction="", questions=["1", "2"]
        )
        ch = ScaleSheet(
            language="ch", title="测试", instruction="", questions=["1", "2", "3"]
        )
        self.assertFalse(paired_shuffle_compatible(en, ch))

    def test_fixed_order_ignores_configured_shuffle_seed(self):
        left = BatchRunner([], RunnerConfig(order_strategy="fixed", random_seed=1))
        right = BatchRunner([], RunnerConfig(order_strategy="fixed", random_seed=999))
        self.assertIsNone(left._effective_order_seed())
        self.assertEqual(left._execution_hash(), right._execution_hash())

    def test_multimodal_content_is_text_plus_image(self):
        content = _user_content("prompt", ["data:image/png;base64,AA=="])
        self.assertIsInstance(content, list)
        self.assertEqual(content[0], {"type": "text", "text": "prompt"})
        self.assertEqual(content[1]["type"], "image_url")

    def test_runner_records_structured_output_and_passes_images(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            image = ScaleImage(anchor="C3", mime_type="image/png", data=PNG_1X1)
            sheet = ScaleSheet(
                language="ch",
                title="困境",
                instruction="情境。",
                questions=["选择 A 或 B", "给出 1 到 7 的评分"],
                images=[image],
            )
            scale = ScaleFile(path=root / "困境.xlsx", name="囚徒困境-评分表", ch=sheet)
            config = RunnerConfig(results_root=root / "results", logs_root=root / "logs")
            config.scale_cfgs[scale.name] = ScaleRunConfig(
                scale_name=scale.name, shuffle_items=False
            )
            runner = BatchRunner([scale], config)
            raw = (
                '{"answers":[{"display_index":1,"answer":"A"},'
                '{"display_index":2,"answer":6}]}'
            )
            with patch("app.runner.call_api", return_value=raw) as mocked_call:
                result_path = runner._one_call(
                    scale, sheet, "ch", 1, "deepseek", "deepseek-v4-pro"
                )

            request_id = mocked_call.call_args.kwargs["client_request_id"]
            self.assertTrue(request_id.isascii())
            request_id.encode("ascii")
            self.assertEqual(len(mocked_call.call_args.kwargs["image_urls"]), 1)
            payload = json.loads(result_path.read_text(encoding="utf-8"))
            self.assertFalse(payload["shuffle_enabled"])
            self.assertIsNone(payload["shuffle_seed"])
            self.assertEqual(payload["shuffle_order"], [1, 2])
            self.assertEqual(payload["items"][0]["answer"], "A")
            self.assertEqual(payload["items"][1]["score"], 6.0)
            self.assertEqual(len(payload["images"]), 1)
            self.assertIn("model_profile", payload)
            self.assertEqual(
                payload["model_profile"]["identity"]["requested_model"],
                "deepseek-v4-pro",
            )

    def test_web_console_displays_readonly_shuffle_status(self):
        client = app.test_client()
        page = client.get("/console").get_data(as_text=True)
        script = (ROOT / "app" / "static" / "js" / "console.js").read_text(encoding="utf-8")
        self.assertIn("题序", page)
        self.assertIn("shuffle-status", script)
        self.assertIn("s.shuffle_status", script)
        self.assertNotIn("row-shuffle", script)

    def test_web_console_removes_shuffle_settings(self):
        page = app.test_client().get("/console").get_data(as_text=True)
        script = (ROOT / "app" / "static" / "js" / "console.js").read_text(encoding="utf-8")
        for obsolete in ("btn-shuffle-all-on", "btn-shuffle-all-off", "btn-apply-order-strategy", 'name="order_strategy"', 'id="random_seed"'):
            self.assertNotIn(obsolete, page)
        self.assertNotIn("setAllShuffle", script)
        self.assertNotIn("selectedOrderStrategy", script)

    def test_web_console_has_one_experiment_scheme_and_derived_condition_status(self):
        client = app.test_client()
        page = client.get("/console").get_data(as_text=True)
        script = (ROOT / "app" / "static" / "js" / "console.js").read_text(encoding="utf-8")
        self.assertIn('id="preset_id" aria-label="实验方案"', page)
        self.assertEqual(page.count('<input id="condition_mode" type="hidden"'), 1)
        for obsolete in ("custom-condition-picker", "condition-mode-badge", "preset-overrides", "preset-warnings"):
            self.assertNotIn(f'id="{obsolete}"', page)
        # Cache-busting revisions are bumped whenever the assets change; assert
        # the assets are referenced rather than freezing a revision number.
        self.assertRegex(page, r"style\.css\?v=\d+")
        self.assertRegex(page, r"console\.js\?v=\d+")
        self.assertIn('id="prompt_contract"', page)
        self.assertNotIn('id="prompt-contract-note"', page)
        self.assertNotIn("PROMPT_CONTRACT_TEXT", script)
        self.assertIn('return activePresetId || "custom";', script)
        self.assertIn("previousValues.forEach", script)
        self.assertIn("updateConditionModeFromPreset", script)

    def test_api_scales_returns_fixed_status_and_image_counts(self):
        image = ScaleImage(anchor="C3", mime_type="image/png", data=PNG_1X1)
        sheet = ScaleSheet(
            language="ch", title="困境", instruction="情境", questions=["题目"], images=[image], shuffle_supported=False
        )
        scale = ScaleFile(path=Path("dilemma.xlsx"), name="电车困境-评分表", ch=sheet)
        with patch("app.webapp.discover_scales", return_value=[scale]), patch(
            "app.webapp.load_user_config", return_value={"scale_vars": {}}
        ):
            data = app.test_client().get("/api/scales?data_dir=./data").get_json()
        item = data["scales"][0]
        self.assertFalse(item["shuffle_items"])
        self.assertEqual(item["shuffle_status"], "固定顺序")
        self.assertEqual(item["ch_images"], 1)
        self.assertIn("ch_architecture", item)
        self.assertEqual(item["ch_architecture"]["kind"], "通用架构")

    def test_api_scales_ignores_obsolete_paired_setting(self):
        en = ScaleSheet(
            language="en", title="Fixture", instruction="", questions=["A", "B"]
        )
        ch = ScaleSheet(
            language="ch", title="测试量表", instruction="", questions=["甲", "乙"]
        )
        mismatch = ScaleFile(
            path=Path("mismatch.xlsx"),
            name="结构不匹配",
            en=en,
            ch=ScaleSheet(language="ch", title="测试量表", instruction="", questions=["甲"]),
        )
        matched = ScaleFile(
            path=Path("matched.xlsx"), name="结构匹配", en=en, ch=ch
        )
        with patch("app.webapp.discover_scales", return_value=[mismatch, matched]), patch(
            "app.webapp.load_user_config",
            return_value={"order_strategy": "paired", "scale_vars": {}},
        ):
            data = app.test_client().get("/api/scales?data_dir=./data").get_json()
        items = {item["name"]: item for item in data["scales"]}
        self.assertTrue(items["结构不匹配"]["shuffle_supported"])
        self.assertTrue(items["结构不匹配"]["shuffle_items"])
        self.assertTrue(items["结构匹配"]["shuffle_supported"])
        self.assertTrue(items["结构匹配"]["shuffle_items"])

    def test_run_uses_profiles_instead_of_obsolete_order_switches(self):
        movable = ScaleSheet(language="ch", title="独立题", instruction="", questions=["甲", "乙"])
        fixed = ScaleSheet(language="ch", title="访谈", instruction="", questions=["前题", "后题"], shuffle_supported=False)
        scales = [
            ScaleFile(path=Path("movable.xlsx"), name="独立题", ch=movable),
            ScaleFile(path=Path("fixed.xlsx"), name="访谈", ch=fixed),
        ]
        cfg = {
            "provider": "deepseek", "model": "fixture-model",
            "order_strategy": "fixed", "capture_network": False,
            "capture_model_catalog": False,
            "scale_vars": {"独立题": {"shuffle_items": False}, "访谈": {"shuffle_items": True}},
        }
        with patch("app.webapp.load_user_config", return_value=cfg), patch(
            "app.webapp.discover_scales", return_value=scales
        ), patch("app.webapp.save_user_config"), patch(
            "app.webapp.proxy_util.apply_proxy_to_env"
        ), patch("app.webapp.threading.Thread"), patch(
            "app.webapp._progress", {"running": False}
        ), patch("app.webapp._runner"), patch("app.webapp._worker"), patch(
            "app.webapp.BatchRunner"
        ) as runner:
            runner.return_value.validate_formal_configuration.return_value = []
            response = app.test_client().post("/api/run/start", json={"order_strategy": "paired"})
            self.assertEqual(response.status_code, 200)
            config = runner.call_args.kwargs["config"]
            self.assertEqual(config.order_strategy, "per_scale")
            self.assertTrue(config.scale_cfgs["独立题"].shuffle_items)
            self.assertFalse(config.scale_cfgs["访谈"].shuffle_items)

    def test_save_ignores_obsolete_order_switches(self):
        cfg = {"order_strategy": "fixed", "scale_vars": {"量表": {"enabled": True}}}
        with patch("app.webapp.load_user_config", return_value=cfg), patch(
            "app.webapp.save_user_config"
        ) as save:
            response = app.test_client().post("/api/config", json={
                "order_strategy": "paired",
                "scale_vars": {"量表": {"enabled": True, "shuffle_items": False}},
            })
            self.assertEqual(response.status_code, 200)
            saved = save.call_args.args[0]
            self.assertEqual(saved["order_strategy"], "per_scale")
            self.assertNotIn("shuffle_items", saved["scale_vars"]["量表"])
            self.assertTrue(saved["scale_vars"]["量表"]["enabled"])

    def test_api_scales_ignores_obsolete_per_scale_switches(self):
        first = ScaleSheet(
            language="ch", title="第一量表", instruction="", questions=["甲"]
        )
        second = ScaleSheet(
            language="ch", title="第二量表", instruction="", questions=["乙"]
        )
        scales = [
            ScaleFile(path=Path("first.xlsx"), name="第一量表", ch=first),
            ScaleFile(path=Path("second.xlsx"), name="第二量表", ch=second),
        ]
        with patch("app.webapp.discover_scales", return_value=scales), patch(
            "app.webapp.load_user_config",
            return_value={
                "order_strategy": "per_scale",
                "scale_vars": {
                    "第一量表": {"shuffle_items": True},
                    "第二量表": {"shuffle_items": False},
                },
            },
        ):
            data = app.test_client().get("/api/scales?data_dir=./data").get_json()
        items = {item["name"]: item for item in data["scales"]}
        self.assertTrue(items["第一量表"]["shuffle_items"])
        self.assertTrue(items["第二量表"]["shuffle_items"])
        self.assertEqual(items["第二量表"]["shuffle_status"], "整行乱序")


if __name__ == "__main__":
    unittest.main()
