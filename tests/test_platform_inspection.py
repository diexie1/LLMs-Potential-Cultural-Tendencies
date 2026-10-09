"""Offline checks for prompts, configuration, retained answers and exports."""

from __future__ import annotations

import json
import random
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from app import config as config_module
from app.experiment_presets import apply_preset
from app.generation_config import simple_platform_generation_config
from app.prompting import build_trial_prompt, cultural_identity_prompt, prompt_preview_sections
from app.runner import BatchRunner, RunnerConfig, ScaleRunConfig
from app.scale_loader import ScaleFile, ScaleSheet
from app.trial_results import EXPORT_NAME, collect_trials, export_trials, read_trial_detail
from app.webapp import app


def fixture_sheet(language="ch"):
    return ScaleSheet(language=language, title="Local metadata", instruction="Original guidance.",
                      questions=["Original question one.", "Original question two."])


def trial(order=1, **overrides):
    return {"run_id": "run_fixture", "scale_name": "Fixture", "language": "ch",
            "provider": "deepseek", "model": "deepseek-v4-pro", "order_id": order,
            "condition_hash": "main", "preset_id": "free_response_v1",
            "generation_config": {"temperature": 0, "top_p": 1}, **overrides}


def write_record(folder, name, value):
    path = folder / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    return path


class PromptInspectionTests(unittest.TestCase):
    def test_identity_modes_only_add_identity_to_main_settings(self):
        main = apply_preset("free_response_v1")
        for name, identity in (("china_identity_v1", "china"), ("usa_identity_v1", "usa")):
            preset = apply_preset(name)
            self.assertEqual(main["generation_config"], preset["generation_config"])
            self.assertEqual(main["execution_config"], preset["execution_config"])
            self.assertEqual(preset["prompt_config"], {**main["prompt_config"], "cultural_identity": identity})

    def test_preview_sections_reassemble_exact_text_in_both_languages(self):
        for language in ("ch", "en"):
            for contract, identity in (("unconstrained", "none"), ("free_scores_only", "none"),
                                       ("unconstrained", "china"), ("unconstrained", "usa")):
                sheet = fixture_sheet(language)
                prompt, _, _ = build_trial_prompt(sheet, language, True, random.Random(7), contract, identity)
                sections = prompt_preview_sections(sheet, language, prompt, contract, identity)
                self.assertEqual("\n\n".join(part["text"] for part in sections), prompt)
                self.assertNotIn(sheet.title, prompt)
                if identity != "none":
                    self.assertTrue(prompt.startswith(cultural_identity_prompt(identity, language) + "\n\n"))

    def test_preview_does_not_consume_global_random_state_or_call_model(self):
        sheet = fixture_sheet()
        scale = ScaleFile(Path("fixture.xlsx"), "Fixture", ch=sheet)
        before = random.getstate()
        with patch("app.webapp.discover_scales", return_value=[scale]), patch("app.runner.call_api") as call:
            response = app.test_client().post("/api/prompt-preview", json={
                "scale_name": scale.name, "language": "ch", "config": {"prompt_config": {"cultural_identity": "china"}}})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(random.getstate(), before)
        call.assert_not_called()
        self.assertIn(sheet.instruction, response.get_json()["prompt"])


class ConfigurationChecks(unittest.TestCase):
    def test_discarded_parameters_cannot_block_simple_configuration(self):
        config = simple_platform_generation_config({"temperature": "0", "top_p": "0.8",
            "top_k": -1, "top_k_mode": "unknown", "thinking_mode": "enabled",
            "extra_body": "invalid legacy field", "response_format": {"type": "json_object"}})
        self.assertEqual(config.temperature, 0)
        self.assertEqual(config.top_p, 0.8)
        self.assertEqual(config.thinking_mode, "disabled")
        self.assertIsNone(config.top_k)
        self.assertIsNone(config.response_format)
        self.assertEqual(config.extra_body, {})

    def test_invalid_temperature_and_top_p_are_rejected(self):
        for value in ({"temperature": -1}, {"temperature": float("nan")}, {"top_p": 2}, [], "bad"):
            with self.subTest(value=value), self.assertRaises((TypeError, ValueError)):
                simple_platform_generation_config(value)

    def test_configuration_writes_are_atomic_and_do_not_modify_input(self):
        with tempfile.TemporaryDirectory(prefix="平台 中文路径 ") as folder:
            path = Path(folder) / "user_config.json"
            with patch.object(config_module, "CONFIG_PATH", path):
                config = {"temperature": 0, "data_dir": str(ROOT / "data")}
                original = dict(config)
                config_module.save_user_config(config)
                self.assertEqual(config, original)
                self.assertEqual(config_module.load_user_config()["temperature"], 0)
                self.assertFalse(list(Path(folder).glob("*.tmp")))
                path.write_text("[]", encoding="utf-8")
                self.assertEqual(config_module.load_user_config(), {})
                path.write_bytes(b"\xff")
                self.assertEqual(config_module.load_user_config(), {})

    def test_external_paths_are_not_rebased_to_a_same_named_project_folder(self):
        with tempfile.TemporaryDirectory(prefix="跨盘 中文路径 ") as folder:
            path = Path(folder).resolve()
            stored = config_module.to_rel_path(path)
            self.assertEqual(config_module.resolve_path(stored), path)

    def test_relocated_platform_loads_from_chinese_path_with_spaces(self):
        with tempfile.TemporaryDirectory(prefix="平台 迁移核对 ") as folder:
            project = Path(folder) / "中文目录 有空格" / "平台"
            project.mkdir(parents=True)
            shutil.copytree(ROOT / "app", project / "app", ignore=shutil.ignore_patterns("__pycache__"))
            shutil.copy2(ROOT / "main.py", project / "main.py")
            code = '''
import json, runpy, sys
from pathlib import Path
target = Path(sys.argv[1]).resolve()
runpy.run_path(str(target / "main.py"), run_name="relocation_check")
from app import config
from app.webapp import app
assert config.ROOT == target and Path.cwd() == target
client = app.test_client()
for route in ("/", "/console", "/chat", "/static/js/console.js", "/static/css/style.css", "/static/bg.jpg"):
    assert client.get(route).status_code == 200, route
config.save_user_config({"results_dir": "./结果 有空格", "data_dir": "./量表 中文"})
assert config.resolve_path(config.load_user_config()["results_dir"]) == target / "结果 有空格"
assert config.CONFIG_PATH.parent == target
print(json.dumps({"relocated": True}))
'''
            completed = subprocess.run([sys.executable, "-c", code, str(project)],
                                       cwd=folder, capture_output=True, text=True,
                                       encoding="utf-8", timeout=30)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue(json.loads(completed.stdout)["relocated"])

    def test_non_object_api_requests_return_json_errors(self):
        for endpoint in ("config", "presets/apply", "prompt-preview", "run/start", "trials/export", "chat/stream"):
            with self.subTest(endpoint=endpoint):
                response = app.test_client().post("/api/" + endpoint, json=[1])
                self.assertEqual(response.status_code, 400)
                self.assertIn("error", response.get_json())

    def test_invalid_counts_do_not_leave_run_marked_active(self):
        scale = ScaleFile(Path("fixture.xlsx"), "Fixture", ch=fixture_sheet())
        for count in (-1, 1.5, "not-a-count"):
            progress = {"running": False}
            with patch("app.webapp.load_user_config", return_value={}), patch("app.webapp.discover_scales", return_value=[scale]), patch("app.webapp._progress", progress):
                response = app.test_client().post("/api/run/start", json={"batch_ch": count})
                self.assertEqual(response.status_code, 400)
                self.assertFalse(progress["running"])

    def test_save_and_start_reject_invalid_execution_settings_without_saving(self):
        cases = [
            {"concurrency": value} for value in (0, 501, 1.5, True, "invalid", float("inf"), float("nan"))
        ] + [
            {"scale_concurrency": value} for value in (0, 501, 2.5, False, "invalid")
        ] + [
            {"batch_ch": -1}, {"batch_en": 1.5}, {"max_retries": 0},
            {"random_seed": 2.5}, {"random_seed": float("inf")},
            {"scale_vars": []}, {"scale_vars": {"Fixture": []}},
            {"scale_vars": {"Fixture": {"en_count": True}}},
        ]
        for endpoint in ("/api/config", "/api/run/start"):
            for data in cases:
                with self.subTest(endpoint=endpoint, data=data), patch(
                    "app.webapp.load_user_config", return_value={}
                ), patch("app.webapp.save_user_config") as save, patch(
                    "app.webapp._progress", {"running": False}
                ), patch("app.webapp.discover_scales") as discover:
                    response = app.test_client().post(endpoint, json=data)
                    self.assertEqual(response.status_code, 400)
                    self.assertIn("error", response.get_json())
                    save.assert_not_called()
                    discover.assert_not_called()

    def test_execution_limits_and_integer_strings_are_saved_without_clamping(self):
        with patch("app.webapp.load_user_config", return_value={}), patch(
            "app.webapp.save_user_config"
        ) as save:
            response = app.test_client().post("/api/config", json={
                "concurrency": "500", "scale_concurrency": 500,
                "batch_en": 0, "batch_ch": "100", "random_seed": "20261009",
                "scale_vars": {"Fixture": {"ch_count": "20", "shuffle_items": False}},
            })
            self.assertEqual(response.status_code, 200)
            saved = save.call_args.args[0]
            self.assertEqual(saved["concurrency"], 500)
            self.assertEqual(saved["scale_concurrency"], 500)
            self.assertEqual(saved["batch_ch"], 100)
            self.assertEqual(saved["random_seed"], 20261009)
            self.assertEqual(saved["scale_vars"]["Fixture"]["ch_count"], 20)
            self.assertNotIn("shuffle_items", saved["scale_vars"]["Fixture"])

    def test_start_rejects_duplicate_runs_and_clears_failed_worker_state(self):
        scale = ScaleFile(Path("fixture.xlsx"), "Fixture", ch=fixture_sheet())
        for fail_thread in (False, True):
            progress = {"running": False}
            cfg = {"provider": "deepseek", "model": "deepseek-v4-pro", "batch_en": 0, "batch_ch": 1}
            with patch("app.webapp.load_user_config", return_value=cfg), patch("app.webapp.save_user_config"), patch("app.webapp.discover_scales", return_value=[scale]), patch("app.webapp._progress", progress), patch("app.webapp._runner"), patch("app.webapp._worker"), patch("app.webapp.proxy_util.apply_proxy_to_env"), patch("app.webapp.BatchRunner") as runner, patch("app.webapp.threading.Thread") as worker:
                runner.return_value.validate_formal_configuration.return_value = []
                if fail_thread:
                    worker.return_value.start.side_effect = RuntimeError("Cannot start thread")
                client = app.test_client()
                first = client.post("/api/run/start", json={})
                self.assertEqual(first.status_code, 400 if fail_thread else 200)
                self.assertEqual(progress["running"], not fail_thread)
                if not fail_thread:
                    self.assertEqual(client.post("/api/run/start", json={}).status_code, 400)
                    self.assertEqual(worker.call_count, 1)


class CompleteAnswerTests(unittest.TestCase):
    def test_results_override_attempts_and_keep_pending_trials(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "plan.jsonl").write_text("\n".join(json.dumps(trial(i)) for i in (1, 2, 3)) + "\n{incomplete", encoding="utf-8")
            write_record(root, "trial_0001_attempt_01.json", trial(attempt_schema_version="attempt-v2", attempt=1, status="error", response_text="old"))
            write_record(root, "trial_0002_attempt_01.json", trial(2, attempt_schema_version="attempt-v2", attempt=1, status="error", response_text="failed"))
            write_record(root, "trial_0002_attempt_02.json", trial(2, attempt_schema_version="attempt-v2", attempt=2, status="unparsed", response_text="latest"))
            write_record(root, "order_0001.json", trial(raw_response="whole paragraph", items=[], response_status="unparsed"))
            rows = collect_trials(root)
            self.assertEqual(len(rows), 3)
            self.assertEqual([r["raw_response"] for r in rows], ["whole paragraph", "latest", ""])
            self.assertEqual([r["source_kind"] for r in rows], ["result", "attempt", "plan"])
            self.assertEqual(rows[2]["status"], "not_recorded")

    def test_different_conditions_remain_separate(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for index, temperature in enumerate((0, 1)):
                value = trial(condition_hash=None, generation_config={"temperature": temperature}, raw_response="answer", items=[])
                write_record(root, f"order_{index}.json", value)
            self.assertEqual(len(collect_trials(root)), 2)

    def test_bad_records_are_skipped_without_hiding_good_answers(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            write_record(root, "order_good.json", trial(raw_response="answer", items=[]))
            for index, order in enumerate((None, "wrong", {}, float("inf"))):
                write_record(root, f"order_bad_{index}.json", trial(order, raw_response="bad", items=[]))
            self.assertEqual(len(collect_trials(root)), 1)

    def test_old_result_without_status_is_shown_as_saved(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            write_record(root, "order_0001.json", trial("0001", raw_response="old answer", items=[]))
            row = collect_trials(root)[0]
            self.assertEqual(row["status"], "legacy_saved")
            self.assertEqual(row["order_id"], 1)
            self.assertEqual(row["raw_response"], "old answer")

    def test_detail_path_is_confined_to_selected_run(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            run = root / "run_fixture"
            run.mkdir()
            write_record(root, "order_outside.json", trial(raw_response="outside", items=[]))
            with self.assertRaises(ValueError):
                read_trial_detail(run, "../order_outside.json", "result")

    def test_excel_preserves_long_unicode_and_never_executes_answer_formulas(self):
        with tempfile.TemporaryDirectory(prefix="导出 中文路径 ") as folder:
            root = Path(folder)
            raw = "原始完整回答\n" + "😀文字" * 18000
            saved = write_record(root, "order_0001.json", trial(raw_response=raw, items=[], response_status="unparsed"))
            original = saved.read_bytes()
            write_record(root, "order_0002.json", trial(2, raw_response='=SUM(1,2)\x00', items=[], response_status="ok"))
            target = export_trials(root)
            self.assertEqual(target.name, EXPORT_NAME)
            self.assertGreater(target.stat().st_size, 0)
            wb = load_workbook(target)
            try:
                sheet = wb.active
                self.assertEqual(sheet.max_row, 3)
                self.assertEqual("".join(cell.value or "" for cell in sheet[2][15:]), raw)
                self.assertEqual(sheet["P3"].value, '=SUM(1,2)\\u0000')
                self.assertEqual(sheet["P3"].data_type, "s")
                self.assertEqual(sheet.freeze_panes, "C2")
            finally:
                wb.close()
            self.assertEqual(saved.read_bytes(), original)

    def test_excel_read_count_does_not_count_unparsed_text_as_success(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            write_record(root, "order_0001.json", trial(raw_response="answer", items=[], response_status="partial", answered_item_count=3, read_item_count=1))
            write_record(root, "order_0002.json", trial(2, raw_response="old answer", items=[], answered_item_count=3))
            target = export_trials(root)
            wb = load_workbook(target)
            try:
                self.assertEqual(wb.active["L2"].value, 1)
                self.assertIsNone(wb.active["L3"].value)
            finally:
                wb.close()

    def test_answer_endpoints_read_records_and_export_workbook(self):
        with tempfile.TemporaryDirectory() as folder, patch("app.webapp.load_user_config", return_value={}):
            root = Path(folder)
            run = root / "run_fixture"
            write_record(run, "order_0001.json", trial(raw_response="complete text", items=[], response_status="unparsed"))
            client = app.test_client()
            query = {"results_dir": str(root), "run": run.name}
            self.assertEqual(client.get("/api/trials/runs", query_string=query).status_code, 200)
            rows = client.get("/api/trials", query_string=query).get_json()["trials"]
            self.assertNotIn("raw_response", rows[0])
            response = client.get("/api/trials/detail", query_string={**query, "source": rows[0]["source"], "kind": rows[0]["source_kind"]})
            self.assertEqual(response.get_json()["raw_response"], "complete text")
            response = client.post("/api/trials/export", json=query)
            self.assertEqual(response.status_code, 200)
            response.close()
            self.assertTrue((run / EXPORT_NAME).exists())

    def test_runner_completion_and_stop_keep_json_csv_and_export(self):
        for stop_after_first in (False, True):
            with self.subTest(stop=stop_after_first), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                scale = ScaleFile(root / "fixture.xlsx", "Fixture", ch=fixture_sheet())
                config = RunnerConfig(results_root=root / "results", logs_root=root / "logs", run_id="run_fixture",
                    concurrency=1, scale_concurrency=1, max_retries=1, capture_network=False, capture_model_catalog=False,
                    scale_cfgs={"Fixture": ScaleRunConfig("Fixture", ch_count=3, en_count=0)})
                runner = BatchRunner([scale], config, log=lambda message: None)
                def response(*args, **kwargs):
                    if stop_after_first:
                        runner.request_stop()
                    return "An unmarked complete response."
                with patch("app.runner.call_api", side_effect=response) as call, patch("app.runner.get_api_credentials", return_value={}), patch("app.runner.proxy_util.detect_system_proxy", return_value=None), patch("app.runner.proxy_util.apply_proxy_to_env"):
                    runner.run()
                run = root / "results" / config.run_id
                self.assertTrue((run / EXPORT_NAME).is_file())
                self.assertEqual(len(list(run.rglob("order_*.json"))), call.call_count)
                self.assertEqual(len(list(run.rglob("order_*.csv"))), call.call_count)
                self.assertEqual(len(collect_trials(run)), 3)


if __name__ == "__main__":
    unittest.main()
