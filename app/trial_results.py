"""Read retained trials and export one complete response per worksheet row."""

from __future__ import annotations

import hashlib
import json
import threading
import uuid
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_NAME = "试次完整回答.xlsx"
_EXPORT_LOCK = threading.Lock()
IDENTITY_NAMES = {"none": "无身份提示", "china": "中国身份", "usa": "美国身份"}
PRESET_NAMES = {
    "free_response_v1": "主分析",
    "free_response_temperature_1_v1": "主分析·温度 1",
    "free_response_scores_only_v1": "主分析·只答分数提示",
    "china_identity_v1": "中国身份条件",
    "usa_identity_v1": "美国身份条件",
}
STATUS_NAMES = {
    "ok": "已保存", "partial": "已保存·部分可逐题提取",
    "unparsed": "已保存·无法逐题提取", "refusal": "已保存·模型拒答",
    "empty_response": "已保存·空回答", "transport_error": "调用失败",
    "not_recorded": "尚无回答记录",
    "legacy_saved": "已保存·旧记录",
}


def read_json(path: Path) -> Dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        return value if isinstance(value, dict) else {}
    except (OSError, UnicodeError, ValueError):
        # A running trial may not have finished writing its JSON yet.
        return {}


def _is_trial(value: Dict[str, Any], kind: str) -> bool:
    if not isinstance(value.get("scale_name"), str) or not value["scale_name"]:
        return False
    if value.get("language") not in ("ch", "en"):
        return False
    try:
        order = value.get("order_id")
        if isinstance(order, bool) or int(order) < 0 or int(order) != float(order):
            return False
    except (TypeError, ValueError, OverflowError):
        return False
    if kind == "plan":
        return True
    if kind == "result":
        return "raw_response" in value and "items" in value
    return "attempt_schema_version" in value and "response_text" in value


def _key(value: Dict[str, Any]) -> tuple:
    generation = value.get("generation_config")
    generation = generation if isinstance(generation, dict) else {}
    condition = value.get("condition_hash") or json.dumps({
        "temperature": generation.get("temperature", value.get("temperature")),
        "prompt_contract": value.get("prompt_contract"),
        "cultural_identity": value.get("cultural_identity"),
    }, sort_keys=True)
    return (
        str(value.get("run_id") or ""), str(value.get("scale_name") or ""),
        str(value.get("language") or ""), str(value.get("provider") or ""),
        str(value.get("model") or value.get("requested_model") or ""),
        str(int(value["order_id"])),
        str(condition),
    )


def normalize_trial(value: Dict[str, Any], source: str, kind: str) -> Dict[str, Any]:
    descriptor = value.get("condition_descriptor")
    descriptor = descriptor if isinstance(descriptor, dict) else {}
    generation = value.get("generation_config")
    generation = generation if isinstance(generation, dict) else {}
    identity = str(value.get("cultural_identity") or descriptor.get("cultural_identity") or "none")
    contract = str(value.get("prompt_contract") or descriptor.get("prompt_contract") or "unconstrained")
    temperature = generation.get("temperature", value.get("temperature"))
    preset_id = str(value.get("preset_id") or "")
    if preset_id and preset_id != "custom":
        experiment = PRESET_NAMES.get(preset_id, preset_id)
    elif identity in {"china", "usa"}:
        experiment = "中国身份条件" if identity == "china" else "美国身份条件"
    elif contract == "free_scores_only":
        experiment = "只答分数提示"
    else:
        experiment = "原指导语"
    if value.get("condition_mode") == "custom":
        experiment = f"自定义（{experiment}）"
    status = str(value.get("response_status") or value.get("status") or (
        "legacy_saved" if kind == "result" else "not_recorded"
    ))
    if kind == "attempt" and status == "error":
        status = "transport_error"
    metadata = value.get("model_identity") or value.get("model_metadata") or {}
    metadata = metadata if isinstance(metadata, dict) else {}
    raw = value.get("raw_response", value.get("response_text", ""))
    return {
        "id": hashlib.sha256(json.dumps(_key(value), ensure_ascii=False).encode()).hexdigest(),
        "run_id": str(value.get("run_id") or ""), "scale_name": value["scale_name"],
        "language": value["language"], "provider": str(value.get("provider") or ""),
        "model": str(value.get("model") or value.get("requested_model") or ""),
        "experiment": experiment, "preset_id": preset_id,
        "cultural_identity": identity, "prompt_contract": contract,
        "temperature": temperature, "top_p": generation.get("top_p"),
        "order_id": int(value["order_id"]), "status": status,
        "status_label": STATUS_NAMES.get(status, status),
        "raw_response": str(raw or ""), "answer_length": len(str(raw or "")),
        "prompt": str(value.get("prompt") or ""),
        "answered_item_count": value.get("answered_item_count"),
        "read_item_count": value.get("read_item_count"),
        "item_count": value.get("item_count", value.get("n_items")),
        "recorded_at": metadata.get("request_completed_at_local") or value.get("recorded_at_utc") or "",
        "source": source, "source_kind": kind,
    }


def collect_trials(folder: Path) -> List[Dict[str, Any]]:
    """Prefer saved results, otherwise retain the last attempt or planned trial."""

    folder = folder.resolve()
    rows: Dict[tuple, Dict[str, Any]] = {}
    for path in sorted(folder.rglob("plan.jsonl")):
        if not path.resolve().is_relative_to(folder):
            continue
        try:
            with path.open(encoding="utf-8-sig") as handle:
                for line in handle:
                    try:
                        value = json.loads(line)
                        if isinstance(value, dict) and _is_trial(value, "plan"):
                            rows[_key(value)] = normalize_trial(value, "", "plan")
                    except (ValueError, TypeError, KeyError):
                        continue
        except (OSError, UnicodeError):
            continue
    latest_attempts: Dict[tuple, int] = {}
    for path in folder.rglob("trial_*_attempt_*.json"):
        if not path.resolve().is_relative_to(folder):
            continue
        value = read_json(path)
        if not _is_trial(value, "attempt"):
            continue
        key = _key(value)
        try:
            attempt = int(value.get("attempt") or 0)
        except (ValueError, TypeError):
            continue
        if attempt < latest_attempts.get(key, -1):
            continue
        base = rows.get(key, {})
        inherited = {
            "preset_id": base.get("preset_id"),
            "generation_config": {"temperature": base.get("temperature"), "top_p": base.get("top_p")},
        }
        latest_attempts[key] = attempt
        rows[key] = normalize_trial({**inherited, **value}, path.relative_to(folder).as_posix(), "attempt")
    for path in folder.rglob("order_*.json"):
        if not path.resolve().is_relative_to(folder):
            continue
        value = read_json(path)
        if _is_trial(value, "result"):
            key = _key(value)
            base = rows.get(key, {})
            rows[key] = normalize_trial(
                {"preset_id": base.get("preset_id"), **value},
                path.relative_to(folder).as_posix(), "result",
            )
    return sorted(rows.values(), key=lambda row: (
        row["run_id"], row["scale_name"], row["language"], row["provider"],
        row["model"], int(row["order_id"] or 0),
    ))


def read_trial_detail(folder: Path, source: str, kind: str) -> Dict[str, Any]:
    if kind not in {"result", "attempt"}:
        raise ValueError("没有可查看的回答文件")
    path = (folder / source).resolve()
    if not path.is_relative_to(folder.resolve()) or path.suffix.lower() != ".json":
        raise ValueError("回答文件路径无效")
    value = read_json(path)
    if not _is_trial(value, kind):
        raise ValueError("回答文件不存在或尚未写入完成")
    return normalize_trial(value, source, kind)


def _excel_text(value: Any) -> str:
    # Keep controls readable instead of silently dropping response characters.
    return ILLEGAL_CHARACTERS_RE.sub(lambda match: f"\\u{ord(match.group()):04x}", str(value or ""))


def _response_parts(text: str) -> List[str]:
    """Keep Unicode characters intact and fit each cell's UTF-16 limit."""

    parts: List[str] = []
    start = units = 0
    for index, char in enumerate(text):
        width = 2 if ord(char) > 0xFFFF else 1
        if units + width > 30000:
            parts.append(text[start:index])
            start, units = index, 0
        units += width
    parts.append(text[start:])
    return parts


def export_trials(folder: Path) -> Path:
    # Automatic completion and an operator's download may coincide.
    with _EXPORT_LOCK:
        return _export_trials(folder)


def _export_trials(folder: Path) -> Path:
    rows = collect_trials(folder)
    wb = Workbook()
    ws = wb.active
    ws.title = "试次完整回答"
    headers = ["运行编号", "量表", "语言", "服务商", "模型", "实验方案", "身份条件", "温度", "Top-p", "试次编号", "运行状态", "逐题提取数", "题目数", "记录时间", "原始记录文件", "完整回答"]
    # Excel permits at most 32,767 characters per cell. Continue in adjacent
    # columns, retaining one row per trial and the original JSON as the source.
    responses = [_response_parts(_excel_text(row["raw_response"])) for row in rows]
    parts = max((len(response) for response in responses), default=1)
    headers += [f"完整回答（续 {index}）" for index in range(2, parts + 1)]
    ws.append(headers)
    for row, response in zip(rows, responses):
        values = [row["run_id"], row["scale_name"], "中文" if row["language"] == "ch" else "英文", row["provider"], row["model"], row["experiment"], IDENTITY_NAMES.get(row["cultural_identity"], row["cultural_identity"]), row["temperature"], row["top_p"], row["order_id"], row["status_label"], row["read_item_count"], row["item_count"], row["recorded_at"], row["source"]]
        values += response + [""] * (parts - len(response))
        ws.append(values)
        ws.row_dimensions[ws.max_row].height = 90
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.value = _excel_text(cell.value)
                cell.data_type = "s"  # Model text must never become a formula.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for index in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 75 if index >= 16 else (30 if index in {1, 2, 5, 15} else 18)
    from openpyxl.comments import Comment
    ws["L1"].comment = Comment("使用 JSON 中的 read_item_count：能读取的分数、选项和开放题文字数量；仅保留了无法解析的逐题文字不计入。旧记录缺少此字段时留空。", "平台")
    ws["P1"].comment = Comment("每行是一个试次。超长回答按列顺序续写，不截断；特殊控制字符显示为 Unicode 转义。原始内容以 JSON 为准。", "平台")
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / EXPORT_NAME
    temporary = folder / f".{uuid.uuid4().hex}.xlsx"
    try:
        wb.save(temporary)
        temporary.replace(target)
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)
    return target
