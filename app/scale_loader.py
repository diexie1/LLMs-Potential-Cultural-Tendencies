"""从 xlsx 加载中英文量表。"""

from __future__ import annotations

import base64
import hashlib
import posixpath
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .config import DATA_DIR


@dataclass
class ScaleImage:
    """Excel worksheet image kept in memory for a multimodal API request."""

    anchor: str
    mime_type: str
    data: bytes = field(repr=False)
    width: Optional[int] = None
    height: Optional[int] = None

    @property
    def data_url(self) -> str:
        encoded = base64.b64encode(self.data).decode("ascii")
        return f"data:{self.mime_type};base64,{encoded}"

    def metadata(self) -> Dict[str, object]:
        return {
            "anchor": self.anchor,
            "mime_type": self.mime_type,
            "bytes": len(self.data),
            "sha256": hashlib.sha256(self.data).hexdigest(),
            "width": self.width,
            "height": self.height,
        }


@dataclass
class ScaleSlot:
    """One explicitly typed response slot in a structured scale profile.

    ``questions`` remains available on :class:`ScaleSheet` for backward
    compatibility, but special profiles use these stable IDs rather than the
    visible numbering embedded in Excel prose.
    """

    slot_id: str
    question: str
    response_type: str = "text"  # text | number | integer | choice | pair_code | ios_pair
    minimum: Optional[float] = None
    maximum: Optional[float] = None
    choices: List[str] = field(default_factory=list)
    context_id: Optional[str] = None
    section_id: Optional[str] = None
    section_label: str = ""
    linked_slot_id: Optional[str] = None
    source_question: Optional[str] = None
    normalization_note: Optional[str] = None
    metadata: Dict[str, object] = field(default_factory=dict)

    def metadata_dict(self) -> Dict[str, object]:
        """Return only serializable, research-relevant slot metadata."""

        out: Dict[str, object] = {
            "slot_id": self.slot_id,
            "response_type": self.response_type,
            "minimum": self.minimum,
            "maximum": self.maximum,
            "choices": list(self.choices),
            "context_id": self.context_id,
            "section_id": self.section_id,
            "section_label": self.section_label,
            "linked_slot_id": self.linked_slot_id,
        }
        if self.source_question and self.source_question != self.question:
            out["source_question"] = self.source_question
        if self.normalization_note:
            out["normalization_note"] = self.normalization_note
        if self.metadata:
            out["metadata"] = dict(self.metadata)
        return out


@dataclass
class ScaleTaskBlock:
    """A contextual task unit that must remain intact in the generated prompt."""

    block_id: str
    label: str
    context: str = ""
    instruction: str = ""
    slots: List[ScaleSlot] = field(default_factory=list)
    # fixed | block | within.  The prompt planner interprets this only when
    # the user explicitly enables shuffling for the scale.
    shuffle_policy: str = "fixed"
    metadata: Dict[str, object] = field(default_factory=dict)


def attach_source_answer_labels(slots: List[ScaleSlot], instruction: str) -> None:
    """Keep explicit source option text as decoder aliases, without guessing labels."""
    marker = re.compile(r"(?<![\w.])([+-]?\d+|[ABX])\s*=\s*|[①②③④⑤⑥⑦⑧⑨⑩]\s*|(?m:^[ \t]*[AB]\s*[.、]\s*)")
    for slot in slots:
        if slot.response_type not in {"integer", "number", "choice"}:
            continue
        aliases: Dict[str, str] = {}
        conflicts = set()
        for source in (instruction, slot.question):
            matches = list(marker.finditer(source))
            if source == instruction:
                codes = {unicodedata.normalize("NFKC", m.group(1) or m.group().strip().rstrip(".、").strip()) for m in matches}
                if slot.choices:
                    if codes != set(slot.choices):
                        continue
                else:
                    numeric_codes = [float(c) for c in codes if re.fullmatch(r"[+-]?\d+", c)]
                    if not numeric_codes or min(numeric_codes) != slot.minimum or max(numeric_codes) != slot.maximum:
                        continue
            for index, match in enumerate(matches):
                raw_code = match.group(1) or match.group().strip().rstrip(".、").strip()
                code = unicodedata.normalize("NFKC", raw_code)
                if slot.choices:
                    if code not in slot.choices:
                        continue
                else:
                    lower = float(slot.minimum) if slot.minimum is not None else float('-inf')
                    upper = float(slot.maximum) if slot.maximum is not None else float('inf')
                    if not re.fullmatch(r"[+-]?\d+", code) or not lower <= float(code) <= upper:
                        continue
                label = source[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(source)]
                # Source definitions end at a newline or a sentence, never at
                # an arbitrary number in surrounding prose.
                label = label.splitlines()[0].strip(" \t,，;；.。") if label.splitlines() else ""
                if not label or len(label) > 100:
                    continue
                for alias in [label, *re.split(r"[/／]", label)]:
                    key = unicodedata.normalize("NFKC", alias).strip().casefold()
                    if key in aliases and aliases[key] != code:
                        conflicts.add(key)
                    aliases[key] = code
        for key in conflicts:
            aliases.pop(key, None)
        slot.metadata.pop("answer_labels", None)
        if aliases:
            slot.metadata["answer_labels"] = aliases


@dataclass
class ScaleSheet:
    language: str  # "en" | "ch"
    title: str
    instruction: str
    questions: List[str] = field(default_factory=list)
    images: List[ScaleImage] = field(default_factory=list)
    profile_id: str = "default_flat_v1"
    profile_version: str = "1"
    profile_label: str = "默认平铺量表"
    profile_error: Optional[str] = None
    blocks: List[ScaleTaskBlock] = field(default_factory=list)
    default_shuffle: Optional[bool] = None
    shuffle_supported: bool = True

    @property
    def n_items(self) -> int:
        if self.blocks:
            return sum(len(block.slots) for block in self.blocks)
        return len(self.questions)

    @property
    def slots(self) -> List[ScaleSlot]:
        """Return response slots in canonical (non-displayed) order."""

        if self.blocks:
            return [slot for block in self.blocks for slot in block.slots]
        return [
            ScaleSlot(slot_id=f"item_{index:03d}", question=question)
            for index, question in enumerate(self.questions, start=1)
        ]

    @property
    def is_profiled(self) -> bool:
        return bool(self.blocks) and self.profile_id != "default_flat_v1"


@dataclass
class ScaleFile:
    path: Path
    name: str
    en: Optional[ScaleSheet] = None
    ch: Optional[ScaleSheet] = None


def paired_shuffle_compatible(en: Optional[ScaleSheet], ch: Optional[ScaleSheet]) -> bool:
    """Return whether English and Chinese sheets share the same shuffle layout.

    A common seed only gives a common permutation when the two language sheets
    expose the same stable response slots and shuffle blocks.  Keep this check
    structural; translated prompt text is expected to differ.
    """

    if en is None or ch is None:
        return False
    if not en.shuffle_supported or not ch.shuffle_supported:
        return False
    if en.profile_id != ch.profile_id or en.n_items != ch.n_items:
        return False
    if [slot.slot_id for slot in en.slots] != [slot.slot_id for slot in ch.slots]:
        return False
    en_blocks = [
        (block.block_id, block.shuffle_policy, tuple(slot.slot_id for slot in block.slots))
        for block in en.blocks
    ]
    ch_blocks = [
        (block.block_id, block.shuffle_policy, tuple(slot.slot_id for slot in block.slots))
        for block in ch.blocks
    ]
    return en_blocks == ch_blocks


def _read_column_a(ws) -> List[str]:
    values: List[str] = []
    for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
        cell = row[0]
        if cell is None:
            values.append("")
        else:
            values.append(str(cell).strip())
    # 去掉尾部空行
    while values and not values[-1]:
        values.pop()
    return values


def _context_from_cell(text: str) -> Optional[str]:
    """Recognize explicit context blocks that must not become scoreable items.

    The old loader classified every non-empty cell from A3 onward as an item.
    To remain compatible with the existing dilemma workbooks, cells beginning
    with common task-setting labels are treated as context. New workbooks
    should use ``[CONTEXT]`` or ``【情境】`` for an unambiguous declaration.
    """

    stripped = text.strip()
    lower = stripped.casefold()
    explicit_markers = (
        "[context]",
        "[instruction]",
        "【情境】",
        "【说明】",
        "【任务设定】",
    )
    for marker in explicit_markers:
        if lower.startswith(marker.casefold()):
            return stripped[len(marker) :].lstrip(" ：:-\n")

    legacy_prefixes = (
        "任务设定",
        "任务说明",
        "作答说明",
        "task setting",
        "task setup",
        "instructions",
        "instruction",
    )
    if lower.startswith(legacy_prefixes):
        return stripped
    return None


def _image_anchor(image, index: int) -> str:
    """Convert an openpyxl image anchor to a stable human-readable cell."""

    anchor = getattr(image, "anchor", None)
    if isinstance(anchor, str):
        return anchor
    start = getattr(anchor, "_from", None)
    if start is None:
        return f"embedded_image_{index}"
    try:
        return f"{get_column_letter(int(start.col) + 1)}{int(start.row) + 1}"
    except (TypeError, ValueError):
        return f"embedded_image_{index}"


def _image_mime_type(image_format: object, data: bytes) -> str:
    """Prefer image signatures because Excel does not always keep ``format``."""

    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if data.startswith(b"BM"):
        return "image/bmp"

    fmt = str(image_format or "").strip().lower()
    return {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "gif": "image/gif",
        "bmp": "image/bmp",
        "webp": "image/webp",
    }.get(fmt, "image/png")


def _read_images(ws) -> List[ScaleImage]:
    """Extract worksheet images without making a single corrupt image fatal."""

    images: List[ScaleImage] = []
    for index, image in enumerate(getattr(ws, "_images", []) or [], start=1):
        try:
            data = image._data()
            if not data:
                continue
            images.append(
                ScaleImage(
                    anchor=_image_anchor(image, index),
                    mime_type=_image_mime_type(getattr(image, "format", None), data),
                    data=data,
                    width=getattr(image, "width", None),
                    height=getattr(image, "height", None),
                )
            )
        except Exception as exc:
            print(f"[warn] Skipping unreadable embedded image #{index}: {exc}")
    return images


_DISPIMG_RE = re.compile(r'DISPIMG\s*\(\s*"([^"]+)"', re.IGNORECASE)


def _local_name(tag: str) -> str:
    """Return an XML local name without assuming a particular vendor prefix."""

    return tag.rsplit("}", 1)[-1]


def _read_cell_images(path: Path, ws) -> List[ScaleImage]:
    """Read Excel/WPS ``DISPIMG`` cell images that openpyxl does not expose.

    Modern Excel/WPS stores these images under ``xl/cellimages.xml`` and
    references them from a ``DISPIMG`` formula. IOS uses this representation,
    so silently relying on ``ws._images`` would lose the actual stimulus.
    """

    formula_ids: Dict[str, str] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if not isinstance(value, str):
                continue
            match = _DISPIMG_RE.search(value)
            if match:
                formula_ids[match.group(1)] = cell.coordinate
    if not formula_ids:
        return []

    try:
        with zipfile.ZipFile(path) as archive:
            if "xl/cellimages.xml" not in archive.namelist():
                return []

            rel_targets: Dict[str, str] = {}
            rel_path = "xl/_rels/cellimages.xml.rels"
            if rel_path in archive.namelist():
                rel_root = ElementTree.fromstring(archive.read(rel_path))
                for rel in rel_root:
                    rel_id = rel.attrib.get("Id")
                    target = rel.attrib.get("Target")
                    if rel_id and target:
                        rel_targets[rel_id] = posixpath.normpath(
                            posixpath.join("xl", target)
                        )

            image_refs: Dict[str, str] = {}
            root = ElementTree.fromstring(archive.read("xl/cellimages.xml"))
            for cell_image in root.iter():
                if _local_name(cell_image.tag) != "cellImage":
                    continue
                image_id = None
                relation_id = None
                for node in cell_image.iter():
                    local = _local_name(node.tag)
                    if local == "cNvPr":
                        image_id = node.attrib.get("name")
                    elif local == "blip":
                        relation_id = next(
                            (
                                value
                                for key, value in node.attrib.items()
                                if key.rsplit("}", 1)[-1] == "embed"
                            ),
                            None,
                        )
                if image_id and relation_id and relation_id in rel_targets:
                    image_refs[image_id] = rel_targets[relation_id]

            images: List[ScaleImage] = []
            for image_id, anchor in formula_ids.items():
                member = image_refs.get(image_id)
                if not member or member not in archive.namelist():
                    continue
                data = archive.read(member)
                if not data:
                    continue
                suffix = Path(member).suffix.lower().lstrip(".")
                images.append(
                    ScaleImage(
                        anchor=anchor,
                        mime_type=_image_mime_type(suffix, data),
                        data=data,
                    )
                )
            return images
    except (OSError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        print(f"[warn] Skipping unreadable DISPIMG stimulus in {path.name}: {exc}")
        return []


def _dedupe_images(images: List[ScaleImage]) -> List[ScaleImage]:
    """Keep one copy of every image payload while preserving first anchors."""

    seen = set()
    out: List[ScaleImage] = []
    for image in images:
        digest = hashlib.sha256(image.data).hexdigest()
        if digest in seen:
            continue
        seen.add(digest)
        out.append(image)
    return out


def _parse_sheet(
    ws,
    language: str,
    source_path: Optional[Path] = None,
) -> Optional[ScaleSheet]:
    cells = _read_column_a(ws)
    if len(cells) < 3:
        return None
    title = cells[0] if cells[0] else "Untitled"
    images = _read_images(ws)

    # A small, explicit registry handles instruments whose Excel cells contain
    # multiple contextual units and response slots.  Import lazily to avoid a
    # module-import cycle: scale_profiles uses the dataclasses declared above.
    from .scale_profiles import ProfileBuildError, build_profile_layout, identify_profile

    profile_id = identify_profile(source_path.stem if source_path else "", title)
    if profile_id:
        if source_path is not None:
            images = _dedupe_images(images + _read_cell_images(source_path, ws))
        try:
            layout = build_profile_layout(profile_id, cells, language, images)
        except ProfileBuildError as exc:
            # Do not silently revert to the flat loader.  That fallback was the
            # source of context-as-item and instruction-as-score errors.
            return ScaleSheet(
                language=language,
                title=title,
                instruction=cells[1] if len(cells) > 1 else "",
                images=images,
                profile_id=profile_id,
                profile_label="专属量表结构（待修复）",
                profile_error=str(exc),
                default_shuffle=False,
                shuffle_supported=False,
            )

        for block in layout.blocks:
            attach_source_answer_labels(block.slots, block.instruction or layout.instruction)
        return ScaleSheet(
            language=language,
            title=title,
            instruction=layout.instruction,
            questions=[slot.question for block in layout.blocks for slot in block.slots],
            images=images,
            profile_id=layout.profile_id,
            profile_version=layout.profile_version,
            profile_label=layout.profile_label,
            blocks=layout.blocks,
            default_shuffle=layout.default_shuffle,
            shuffle_supported=layout.shuffle_supported,
        )

    instruction_blocks = [cells[1]] if len(cells) > 1 and cells[1] else []
    questions: List[str] = []
    for cell in cells[2:]:
        if not cell:
            continue
        context = _context_from_cell(cell)
        if context is not None:
            if context:
                instruction_blocks.append(context)
        else:
            questions.append(cell)
    if not questions:
        return None
    return ScaleSheet(
        language=language,
        title=title,
        instruction="\n\n".join(instruction_blocks),
        questions=questions,
        images=images,
    )


def load_scale_file(path: Path) -> ScaleFile:
    # ``read_only`` worksheets do not expose embedded images.
    # Keep formula text as well: Excel/WPS DISPIMG formulas identify cell-bound
    # stimuli that openpyxl does not expose through ``ws._images``.
    wb = load_workbook(path, read_only=False, data_only=False)
    try:
        name = path.stem
        en = ch = None
        # Sheet1 = 英文, Sheet2 = 中文
        if "Sheet1" in wb.sheetnames:
            en = _parse_sheet(wb["Sheet1"], "en", path)
        if "Sheet2" in wb.sheetnames:
            ch = _parse_sheet(wb["Sheet2"], "ch", path)
        # 兜底：按顺序取前两张
        if en is None and wb.sheetnames:
            en = _parse_sheet(wb[wb.sheetnames[0]], "en", path)
        if ch is None and len(wb.sheetnames) > 1:
            ch = _parse_sheet(wb[wb.sheetnames[1]], "ch", path)
        return ScaleFile(path=path, name=name, en=en, ch=ch)
    finally:
        wb.close()


def discover_scales(data_dir: Optional[Path] = None) -> List[ScaleFile]:
    data_dir = Path(data_dir) if data_dir else DATA_DIR
    if not data_dir.exists():
        return []
    scales: List[ScaleFile] = []
    for path in sorted(data_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            scales.append(load_scale_file(path))
        except Exception as exc:
            print(f"[warn] 跳过无法读取的文件 {path.name}: {exc}")
    return scales
